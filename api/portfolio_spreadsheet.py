"""Bounded spreadsheet reading and explicit broker column mapping. No formula evaluation."""
import csv
import io
import json
import re
import zipfile
from datetime import date, datetime
from api.family_portfolio_service import DETAIL_NUMBERS, DETAIL_TEXT, PortfolioError, parse_csv, require

FIELDS = {'date', 'symbol', 'quantity', 'amount', 'currency', 'type', 'externalId'} | DETAIL_NUMBERS | DETAIL_TEXT
BROKERS = {'toss', 'kiwoom', 'korea-investment', 'miraeasset', 'samsung', 'other'}
TYPE_NAMES = {'매수': 'buy', '매도': 'sell', '배당': 'dividend', '배당금': 'dividend', '배당취소': 'dividend_reversal',
              '입금': 'cash_in', '출금': 'cash_out', '입고': 'transfer_in', '출고': 'transfer_out', '분할': 'split'}


def _cell(cell):
    require(cell.data_type not in {'f', 'e'}, '수식·오류 셀은 지원하지 않습니다. 값만 복사한 XLSX를 사용하세요.')
    value = cell.value
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return str(value)
    require(not isinstance(value, (int, float)) or abs(value) < 1e15,
            '15자리 이상 숫자는 Excel에서 정밀도가 손실될 수 있습니다. 원본에서 텍스트로 내보내세요.')
    if isinstance(value, (int, float)) and re.fullmatch(r'0{2,12}', cell.number_format or '') and value == int(value):
        return str(int(value)).zfill(len(cell.number_format))
    return '' if value is None else str(value)


def spreadsheet_table(content, filename, sheet='', header_row=1, encoding='utf-8-sig'):
    require(isinstance(header_row, int) and 1 <= header_row <= 20, '제목 행은 1~20행으로 선택하세요.')
    sheets = []
    if filename.lower().endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                require(len(members) <= 1000 and sum(m.file_size for m in members) <= 20 * 1024 * 1024,
                        '압축 해제 크기가 너무 큰 XLSX입니다.')
                require(not any('vbaproject' in m.filename.lower() for m in members), '매크로 파일은 지원하지 않습니다.')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheets = workbook.sheetnames
                require(0 < len(sheets) <= 20, '워크시트는 최대 20개까지 지원합니다.')
                selected = sheet or sheets[0]
                require(selected in sheets, '워크시트를 확인하세요.')
                page = workbook[selected]
                require((page.max_column or 0) <= 50 and (page.max_row or 0) <= header_row + 200,
                        '워크시트는 제목 포함 50열·자료 200행까지 지원합니다. 불필요한 행·열을 제거하세요.')
                raw = [[_cell(cell) for cell in row] for row in page.iter_rows(max_row=header_row + 201, max_col=min(51, page.max_column or 50))]
            finally:
                workbook.close()
        except PortfolioError:
            raise
        except Exception:
            raise PortfolioError('XLSX 파일을 읽지 못했습니다. 암호·수식·파일 형식을 확인하세요.')
    else:
        require(filename.lower().endswith('.csv'), 'CSV 또는 XLSX만 지원합니다. XLS·스캔 PDF는 지원하지 않습니다.')
        require(encoding in {'utf-8-sig', 'cp949'}, 'CSV 인코딩을 선택하세요.')
        try:
            reader = csv.reader(io.StringIO(content.decode(encoding)), strict=True)
            raw = []
            for row in reader:
                raw.append(row)
                require(len(row) <= 50 and len(raw) <= header_row + 200, '자료는 50열·200행까지 지원합니다.')
        except (UnicodeError, csv.Error):
            raise PortfolioError('CSV 인코딩 또는 열 형식을 확인하세요.')
        selected = ''
    require(len(raw) >= header_row, '제목 행을 찾지 못했습니다.')
    headers = [str(v).strip()[:200] for v in raw[header_row - 1]]
    while headers and not headers[-1]:
        headers.pop()
    require(0 < len(headers) <= 50 and any(headers), '열 제목을 확인하세요.')
    rows = [row for row in raw[header_row:] if any(str(v).strip() for v in row)]
    require(0 < len(rows) <= 200, '자료는 1~200행까지 지원합니다.')
    require(all(not any(str(v).strip() for v in row[len(headers):]) for row in rows), '제목 없는 열에 자료가 있습니다.')
    return {'sheets': sheets, 'sheet': selected, 'headers': headers, 'rows': rows}


def inspect_spreadsheet(content, filename, form):
    table = spreadsheet_table(content, filename, form.get('sheet', ''), _header_row(form), form.get('encoding', 'utf-8-sig'))
    return {key: table[key] for key in ('sheets', 'sheet', 'headers')} | {'rowCount': len(table['rows'])}


def _header_row(form):
    try:
        return int(form.get('headerRow', 1))
    except (TypeError, ValueError):
        raise PortfolioError('제목 행을 확인하세요.')


def mapped_spreadsheet(content, filename, form, kind, as_of):
    table = spreadsheet_table(content, filename, form.get('sheet', ''), _header_row(form), form.get('encoding', 'utf-8-sig'))
    try:
        mapping = json.loads(form.get('mapping', '{}'))
    except (ValueError, TypeError):
        raise PortfolioError('열 연결을 확인하세요.')
    require(isinstance(mapping, dict) and set(mapping) <= FIELDS, '지원하지 않는 열 연결입니다.')
    for key, index in mapping.items():
        require(type(index) is int and 0 <= index < len(table['headers']), '열 연결 범위를 확인하세요.')
    require(len(set(mapping.values())) == len(mapping), '같은 열을 여러 항목에 연결할 수 없습니다.')
    required = {'symbol', 'quantity'} if kind == 'holdings' else {'date', 'symbol', 'type', 'externalId'}
    require(required <= set(mapping), '필수 열을 모두 연결하세요. 거래 고유번호가 없는 자료는 표준 CSV로 검토해 주세요.')
    default_currency = form.get('defaultCurrency', '')
    require('currency' in mapping or default_currency in {'USD', 'KRW'}, '통화 열 또는 기본 통화를 확인하세요.')
    broker = form.get('broker', 'other')
    require(broker in BROKERS, '증권사 구분을 확인하세요.')
    result = []
    for raw in table['rows']:
        row = {key: str(raw[index]).strip() if index < len(raw) else '' for key, index in mapping.items()}
        row.setdefault('currency', default_currency)
        if kind == 'holdings':
            row.setdefault('date', as_of)
        if row.get('date'):
            value = row['date'].strip()
            if re.fullmatch(r'\d{8}', value):
                value = f'{value[:4]}-{value[4:6]}-{value[6:]}'
            row['date'] = re.sub(r'[./]', '-', value)
        for key in {'quantity', 'amount'} | DETAIL_NUMBERS:
            if key in row:
                row[key] = row[key].replace(',', '').strip()
        if kind == 'transactions':
            row['rawType'] = row.get('rawType') or row['type']
            row['type'] = TYPE_NAMES.get(row['type'], row['type'])
        result.append(row)
    return result, f'broker-mapped-v1:{broker}'


def standard_xlsx(content, filename):
    table = spreadsheet_table(content, filename)
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(table['headers'])
    for row in table['rows']:
        writer.writerow(row[:len(table['headers'])] + [''] * max(0, len(table['headers']) - len(row)))
    return parse_csv(stream.getvalue().encode('utf-8'))
